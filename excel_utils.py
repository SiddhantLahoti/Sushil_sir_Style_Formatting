from copy import copy
from openpyxl.formula.translate import Translator
from openpyxl.utils import coordinate_to_tuple


def get_actual_max_row(ws):
    """Finds the last row containing actual data, ignoring empty formatted rows."""
    for r in range(ws.max_row, 0, -1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None and str(val).strip() != "":
                return r
    return 1


def copy_cell(source_cell, target_cell, row_delta):
    """Copies cell value, translated formulas, and visual styling."""
    if source_cell.data_type == "f" or (
        isinstance(source_cell.value, str) and source_cell.value.startswith("=")
    ):
        try:
            target_cell.value = Translator(
                source_cell.value, origin=source_cell.coordinate
            ).translate_formula(row_delta=row_delta)
        except Exception:
            target_cell.value = source_cell.value
    else:
        target_cell.value = source_cell.value

    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = copy(source_cell.number_format)
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)


def find_style_target_cell(ws, max_row, max_col, manual_target=None):
    """Finds the cell coordinate containing the Style Number."""
    if manual_target:
        return coordinate_to_tuple(manual_target)

    keywords = ["style no", "style no.", "style #", "style_no", "style code", "style"]

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            val = str(ws.cell(row=r, column=c).value or "").strip().lower()
            clean_val = val.replace(":", "").replace(".", "").replace("#", "").strip()

            if clean_val in keywords or any(kw in val for kw in keywords):
                for offset in (1, 2):
                    adj_cell = ws.cell(row=r, column=c + offset)
                    adj_val = str(adj_cell.value or "").strip().lower()
                    if adj_cell.value is not None and not any(
                        lbl in adj_val for lbl in ["buyer", "date", "season", "order", "qty"]
                    ):
                        return (r, c + offset)
                return (r + 1, c)

    raise ValueError("Could not automatically locate the Style No cell in the template.")


def find_template_diamond_columns(ws, start_row, end_row, max_col, header_aliases):
    """
    Finds the header row and column indices for diamond data inside the template block.
    Returns (header_row, col_map) where col_map is e.g. {'sieve': 3, 'qty': 7}.
    """
    for r in range(start_row, end_row + 1):
        found_cols = {}
        for c in range(1, max_col + 1):
            cell_val = str(ws.cell(row=r, column=c).value or "").strip().lower()
            if not cell_val:
                continue
            for key, aliases in header_aliases.items():
                if cell_val in aliases or any(alias in cell_val for alias in aliases):
                    if key not in found_cols:
                        found_cols[key] = c

        # Found the diamond table header row if at least 3 matching columns are detected
        if len(found_cols) >= 3:
            return r, found_cols

    raise ValueError("Could not locate Diamond Table headers ('SIEVE', 'QTY', etc.) in the template.")