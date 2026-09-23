import os
from copy import copy
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import coordinate_to_tuple, get_column_letter

# ==============================================================================
# CONFIGURATION
# ==============================================================================
STYLE_BREAKUP_PATH = "Style Breakup.xlsx"
COSTING_TEMPLATE_PATH = "Costing-Format_1_product.xlsx"
OUTPUT_PATH = "Costing_Breakup_Generated.xlsx"

# Set to the last row number of your common header (copied only ONCE at the top).
# For example, if rows 1 to 3 are company/title headers, set HEADER_LAST_ROW = 3.
# If set to None, it will automatically start product blocks at the Style No row.
HEADER_LAST_ROW = 4

# Set the exact cell of the Style No in the template (e.g., "B2" or "B3").
# If None, the script will automatically detect the label and replace its value.
STYLE_TARGET_CELL = "B6"

# Blank rows between consecutive product blocks
SPACER_ROWS = 2


def get_products_from_breakup(file_path):
    """Reads style numbers from rows where 'Cust' is populated."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    header_row_idx = None
    cust_col_idx = None
    style_col_idx = None

    for row in ws.iter_rows(values_only=False):
        for cell in row:
            val = str(cell.value).strip().lower() if cell.value else ""
            if val == "cust":
                cust_col_idx = cell.column
                header_row_idx = cell.row
            elif val in ["style no", "style no.", "style_no", "style", "style code", "style #"]:
                style_col_idx = cell.column

        if header_row_idx and cust_col_idx and style_col_idx:
            break

    if not header_row_idx or not style_col_idx or not cust_col_idx:
        raise ValueError("Could not find 'Cust' and 'Style No' column headers in Style Breakup.")

    products = []
    for r in range(header_row_idx + 1, ws.max_row + 1):
        cust_val = ws.cell(row=r, column=cust_col_idx).value
        if cust_val is not None and str(cust_val).strip() != "":
            style_no = ws.cell(row=r, column=style_col_idx).value
            if style_no is not None and str(style_no).strip() != "":
                products.append(str(style_no).strip())

    wb.close()
    return products


def get_actual_max_row(ws):
    """Finds the last row containing actual data, ignoring empty formatted rows."""
    for r in range(ws.max_row, 0, -1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None and str(val).strip() != "":
                return r
    return 1


def find_style_target_cell(ws, max_row, max_col):
    """
    Finds the exact cell containing the style number to be replaced.
    Handles colons, merged spans, and raises an error if not found.
    """
    if STYLE_TARGET_CELL:
        return coordinate_to_tuple(STYLE_TARGET_CELL)

    keywords = ["style no", "style no.", "style #", "style_no", "style code", "style"]

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            val = str(ws.cell(row=r, column=c).value or "").strip().lower()
            clean_val = val.replace(":", "").replace(".", "").replace("#", "").strip()

            if clean_val in keywords or any(kw in val for kw in keywords):
                # Check adjacent cells up to 2 columns to account for merged label cells
                for offset in (1, 2):
                    adj_cell = ws.cell(row=r, column=c + offset)
                    adj_val = str(adj_cell.value or "").strip().lower()
                    if adj_cell.value is not None and not any(lbl in adj_val for lbl in ["buyer", "date", "season", "order", "qty", "item"]):
                        return (r, c + offset)

                # Fallback to the row immediately beneath the label
                return (r + 1, c)

    raise ValueError(
        "Could not automatically detect the Style No cell in the template. "
        "Please set STYLE_TARGET_CELL in CONFIGURATION (e.g., STYLE_TARGET_CELL = 'B3')."
    )


def copy_cell(source_cell, target_cell, row_delta):
    """Copies cell value, translated formulas, and styles."""
    if source_cell.data_type == "f" or (
        isinstance(source_cell.value, str) and source_cell.value.startswith("=")
    ):
        try:
            target_cell.value = Translator(
                source_cell.value, origin=source_cell.coordinate
            ).translate_formula(row_delta=row_delta)
        except Exception:
            target_cell.value = source_cell.value
    else:
        target_cell.value = source_cell.value

    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = copy(source_cell.number_format)
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)


def generate_costing_sheet():
    products = get_products_from_breakup(STYLE_BREAKUP_PATH)
    print(f"Loaded {len(products)} products from {STYLE_BREAKUP_PATH}")

    wb_template = openpyxl.load_workbook(COSTING_TEMPLATE_PATH)
    ws_template = wb_template.active

    # Determine real boundaries
    tpl_actual_max_row = get_actual_max_row(ws_template)
    tpl_max_col = ws_template.max_column
    tpl_merges = list(ws_template.merged_cells.ranges)

    # Locate the cell where the style number resides
    target_r_rel, target_c_rel = find_style_target_cell(ws_template, tpl_actual_max_row, tpl_max_col)
    print(f"Template Style No located at: Row {target_r_rel}, Column {target_c_rel}")

    # Determine header vs product block rows
    header_end = HEADER_LAST_ROW if HEADER_LAST_ROW is not None else max(0, target_r_rel - 1)
    product_start = header_end + 1
    product_block_height = tpl_actual_max_row - product_start + 1

    print(f"Header range (copied once): Rows 1 to {header_end}")
    print(f"Product block range: Rows {product_start} to {tpl_actual_max_row} (Height: {product_block_height} rows)")

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = ws_template.title

    # Copy column widths
    for col_idx in range(1, tpl_max_col + 1):
        col_letter = get_column_letter(col_idx)
        if col_letter in ws_template.column_dimensions:
            ws_out.column_dimensions[col_letter].width = ws_template.column_dimensions[col_letter].width

    current_out_row = 1

    # --------------------------------------------------------------------------
    # 1. Copy Header Section (Executed ONCE)
    # --------------------------------------------------------------------------
    if header_end > 0:
        for r in range(1, header_end + 1):
            if r in ws_template.row_dimensions:
                ws_out.row_dimensions[r].height = ws_template.row_dimensions[r].height

            for c in range(1, tpl_max_col + 1):
                src_cell = ws_template.cell(row=r, column=c)
                tgt_cell = ws_out.cell(row=r, column=c)
                copy_cell(src_cell, tgt_cell, row_delta=0)

        # Copy merged cells belonging to the header
        for merge in tpl_merges:
            if merge.max_row <= header_end:
                ws_out.merge_cells(
                    start_row=merge.min_row,
                    end_row=merge.max_row,
                    start_column=merge.min_col,
                    end_column=merge.max_col,
                )

        current_out_row = header_end + 1

    # --------------------------------------------------------------------------
    # 2. Copy Product Blocks
    # --------------------------------------------------------------------------
    for idx, style_num in enumerate(products):
        # Add 2 spacer lines before subsequent products
        if idx > 0:
            current_out_row += SPACER_ROWS

        block_target_start = current_out_row
        row_delta = block_target_start - product_start

        # Copy product block rows and translate formulas
        for src_r in range(product_start, tpl_actual_max_row + 1):
            tgt_r = src_r + row_delta
            if src_r in ws_template.row_dimensions:
                ws_out.row_dimensions[tgt_r].height = ws_template.row_dimensions[src_r].height

            for c in range(1, tpl_max_col + 1):
                src_cell = ws_template.cell(row=src_r, column=c)
                tgt_cell = ws_out.cell(row=tgt_r, column=c)
                copy_cell(src_cell, tgt_cell, row_delta=row_delta)

        # Replicate merged cells belonging to the product block
        for merge in tpl_merges:
            if merge.min_row >= product_start and merge.max_row <= tpl_actual_max_row:
                ws_out.merge_cells(
                    start_row=merge.min_row + row_delta,
                    end_row=merge.max_row + row_delta,
                    start_column=merge.min_col,
                    end_column=merge.max_col,
                )

        # Overwrite the style number
        actual_style_row = block_target_start + (target_r_rel - product_start)
        ws_out.cell(row=actual_style_row, column=target_c_rel).value = style_num

        print(f"Product {idx + 1}/{len(products)} written (Rows {block_target_start} to {block_target_start + product_block_height - 1}) -> Style: {style_num}")

        current_out_row += product_block_height

    wb_out.save(OUTPUT_PATH)
    print(f"\nSuccessfully generated {OUTPUT_PATH}")


if __name__ == "__main__":
    generate_costing_sheet()